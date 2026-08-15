"""Generate native DOCX and PDF files from an AI Markdown report."""

from __future__ import annotations

import html
import re
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


REPORT_TITLE = "C2Sherlock AI分析报告"
ACCENT = "064E3B"

WORD_BODY_CJK_FONT = "SimSun"
WORD_HEADING_CJK_FONT = "Microsoft YaHei"
WORD_LATIN_FONT = "Times New Roman"
WORD_BODY_SIZE = 12
WORD_HEADING_SIZE = 14

FONT_DIR = Path(__file__).resolve().parent.parent / "assets" / "fonts"
PDF_FONT_FILES = {
    "C2SourceHanSansRegular": "SourceHanSansSC-Regular.ttf",
    "C2SourceHanSansBold": "SourceHanSansSC-Bold.ttf",
    "C2LiberationSerifRegular": "LiberationSerif-Regular.ttf",
    "C2LiberationSerifBold": "LiberationSerif-Bold.ttf",
}


class ReportFontError(RuntimeError):
    """Raised when the fonts required for an embedded PDF are unavailable."""


def _clean_inline(text: str) -> str:
    text = re.sub(r"!\[[^\]]*]\([^)]*\)", "", text)
    text = re.sub(r"\[([^\]]+)]\([^)]*\)", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    return text.replace("**", "").replace("__", "").strip()


def _is_latin_pdf_character(character: str) -> bool:
    # ASCII letters, digits, spaces and punctuation use Liberation Serif.
    # CJK characters and full-width punctuation use Source Han Sans.
    return ord(character) < 128


def _pdf_font_span(text: str, *, bold: bool) -> str:
    if not text:
        return ""
    cjk_font = "C2SourceHanSansBold" if bold else "C2SourceHanSansRegular"
    latin_font = "C2LiberationSerifBold" if bold else "C2LiberationSerifRegular"
    groups: list[str] = []
    start = 0
    current_is_latin = _is_latin_pdf_character(text[0])
    for index, character in enumerate(text[1:], start=1):
        is_latin = _is_latin_pdf_character(character)
        if is_latin != current_is_latin:
            font_name = latin_font if current_is_latin else cjk_font
            groups.append(f'<font name="{font_name}">{html.escape(text[start:index], quote=False)}</font>')
            start = index
            current_is_latin = is_latin
    font_name = latin_font if current_is_latin else cjk_font
    groups.append(f'<font name="{font_name}">{html.escape(text[start:], quote=False)}</font>')
    return "".join(groups)


def _inline_pdf(text: str, *, force_bold: bool = False) -> str:
    text = re.sub(r"!\[[^\]]*]\([^)]*\)", "", text)
    text = re.sub(r"\[([^\]]+)]\([^)]*\)", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    if force_bold:
        return _pdf_font_span(text.replace("**", "").replace("__", ""), bold=True)

    result: list[str] = []
    bold = False
    cursor = 0
    for marker in re.finditer(r"\*\*|__", text):
        result.append(_pdf_font_span(text[cursor:marker.start()], bold=bold))
        bold = not bold
        cursor = marker.end()
    result.append(_pdf_font_span(text[cursor:], bold=bold))
    return "".join(result)


def parse_markdown(markdown: str) -> list[tuple[str, object]]:
    blocks: list[tuple[str, object]] = []
    paragraph: list[str] = []
    lines = markdown.replace("\r\n", "\n").split("\n")

    def flush_paragraph() -> None:
        if paragraph:
            content = " ".join(item.strip() for item in paragraph).strip()
            if content:
                blocks.append(("paragraph", content))
            paragraph.clear()

    index = 0
    while index < len(lines):
        raw = lines[index].strip()
        if not raw:
            flush_paragraph()
            index += 1
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", raw)
        if heading:
            flush_paragraph()
            blocks.append((f"h{min(len(heading.group(1)), 3)}", heading.group(2).strip()))
            index += 1
            continue
        bullet = re.match(r"^[-*+]\s+(.+)$", raw)
        numbered = re.match(r"^\d+[.)]\s+(.+)$", raw)
        if bullet or numbered:
            flush_paragraph()
            kind = "bullet" if bullet else "numbered"
            matcher = r"^[-*+]\s+(.+)$" if kind == "bullet" else r"^\d+[.)]\s+(.+)$"
            items: list[str] = []
            while index < len(lines):
                match = re.match(matcher, lines[index].strip())
                if not match:
                    break
                items.append(match.group(1).strip())
                index += 1
            blocks.append((kind, items))
            continue
        if raw.startswith("|") and raw.endswith("|"):
            flush_paragraph()
            rows: list[list[str]] = []
            while index < len(lines):
                candidate = lines[index].strip()
                if not (candidate.startswith("|") and candidate.endswith("|")):
                    break
                cells = [cell.strip() for cell in candidate.strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
                    rows.append(cells)
                index += 1
            if rows:
                blocks.append(("table", rows))
            continue
        paragraph.append(raw)
        index += 1
    flush_paragraph()
    return blocks


def _set_run_font(
    run,
    size: float,
    *,
    cjk_font: str,
    bold: bool = False,
    color: str = "172033",
) -> None:
    run.font.name = WORD_LATIN_FONT
    run_fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    run_fonts.set(qn("w:ascii"), WORD_LATIN_FONT)
    run_fonts.set(qn("w:hAnsi"), WORD_LATIN_FONT)
    run_fonts.set(qn("w:cs"), WORD_LATIN_FONT)
    run_fonts.set(qn("w:eastAsia"), cjk_font)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def _format_docx_paragraph(paragraph, *, before: float = 0, after: float = 0) -> None:
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.5


def build_docx(markdown: str) -> bytes:
    document = Document()
    now = datetime.now(timezone.utc)
    document.core_properties.created = now
    document.core_properties.modified = now
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _format_docx_paragraph(title, after=18)
    _set_run_font(
        title.add_run(REPORT_TITLE),
        WORD_HEADING_SIZE,
        cjk_font=WORD_HEADING_CJK_FONT,
        bold=True,
        color=ACCENT,
    )
    for kind, value in parse_markdown(markdown):
        if kind.startswith("h"):
            level = int(kind[1])
            paragraph = document.add_paragraph()
            _format_docx_paragraph(paragraph, before=12 if level == 1 else 8, after=6)
            _set_run_font(
                paragraph.add_run(_clean_inline(str(value))),
                WORD_HEADING_SIZE,
                cjk_font=WORD_HEADING_CJK_FONT,
                bold=True,
                color=ACCENT,
            )
        elif kind == "paragraph":
            paragraph = document.add_paragraph()
            _format_docx_paragraph(paragraph, after=8)
            _set_run_font(
                paragraph.add_run(_clean_inline(str(value))),
                WORD_BODY_SIZE,
                cjk_font=WORD_BODY_CJK_FONT,
            )
        elif kind in {"bullet", "numbered"}:
            style = "List Bullet" if kind == "bullet" else "List Number"
            for item in value:  # type: ignore[union-attr]
                paragraph = document.add_paragraph(style=style)
                _format_docx_paragraph(paragraph, after=4)
                _set_run_font(
                    paragraph.add_run(_clean_inline(str(item))),
                    WORD_BODY_SIZE,
                    cjk_font=WORD_BODY_CJK_FONT,
                )
        elif kind == "table":
            rows = value  # type: ignore[assignment]
            columns = max(len(row) for row in rows)
            table = document.add_table(rows=0, cols=columns)
            table.style = "Table Grid"
            for row_index, row in enumerate(rows):
                cells = table.add_row().cells
                for column_index in range(columns):
                    cell = cells[column_index]
                    cell.text = ""
                    text = row[column_index] if column_index < len(row) else ""
                    _format_docx_paragraph(cell.paragraphs[0])
                    _set_run_font(
                        cell.paragraphs[0].add_run(_clean_inline(text)),
                        WORD_BODY_SIZE,
                        cjk_font=WORD_BODY_CJK_FONT,
                        bold=row_index == 0,
                    )
            document.add_paragraph().paragraph_format.space_after = Pt(3)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _register_pdf_fonts() -> None:
    missing = [filename for filename in PDF_FONT_FILES.values() if not (FONT_DIR / filename).is_file()]
    if missing:
        names = "、".join(missing)
        raise ReportFontError(f"PDF 字体文件缺失：{names}。请将字体文件放入 {FONT_DIR}")

    for font_name, filename in PDF_FONT_FILES.items():
        if font_name in pdfmetrics.getRegisteredFontNames():
            continue
        font_path = FONT_DIR / filename
        try:
            pdfmetrics.registerFont(TTFont(font_name, str(font_path)))
        except Exception as exc:
            raise ReportFontError(
                f"PDF 字体无法加载：{font_path}。请确认文件完整且为 ReportLab 支持的 TTF/OTF 字体。"
            ) from exc


def build_pdf(markdown: str) -> bytes:
    _register_pdf_fonts()
    output = BytesIO()
    doc = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=16 * mm,
        leftMargin=16 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title=REPORT_TITLE,
        author="C2Sherlock",
    )
    base = getSampleStyleSheet()
    body = ParagraphStyle(
        "ReportBody",
        parent=base["BodyText"],
        fontName="C2SourceHanSansRegular",
        fontSize=12,
        leading=18,
        spaceAfter=8,
        textColor=colors.HexColor("#172033"),
    )
    title_style = ParagraphStyle(
        "ReportTitle",
        parent=base["Title"],
        fontName="C2SourceHanSansBold",
        fontSize=14,
        leading=21,
        alignment=TA_CENTER,
        spaceAfter=18,
        textColor=colors.HexColor(f"#{ACCENT}"),
    )
    heading_styles = {
        "h1": ParagraphStyle("ReportH1", parent=base["Heading1"], fontName="C2SourceHanSansBold", fontSize=14, leading=21, spaceBefore=12, spaceAfter=6, textColor=colors.HexColor(f"#{ACCENT}")),
        "h2": ParagraphStyle("ReportH2", parent=base["Heading2"], fontName="C2SourceHanSansBold", fontSize=14, leading=21, spaceBefore=9, spaceAfter=5, textColor=colors.HexColor(f"#{ACCENT}")),
        "h3": ParagraphStyle("ReportH3", parent=base["Heading3"], fontName="C2SourceHanSansBold", fontSize=14, leading=21, spaceBefore=7, spaceAfter=4, textColor=colors.HexColor(f"#{ACCENT}")),
    }
    story = [Paragraph(_inline_pdf(REPORT_TITLE, force_bold=True), title_style)]
    for kind, value in parse_markdown(markdown):
        if kind.startswith("h"):
            story.append(Paragraph(_inline_pdf(str(value), force_bold=True), heading_styles[kind]))
        elif kind == "paragraph":
            story.append(Paragraph(_inline_pdf(str(value)), body))
        elif kind in {"bullet", "numbered"}:
            items = [ListItem(Paragraph(_inline_pdf(str(item)), body)) for item in value]  # type: ignore[union-attr]
            story.append(ListFlowable(items, bulletType="bullet" if kind == "bullet" else "1", leftIndent=16))
            story.append(Spacer(1, 4))
        elif kind == "table":
            rows = value  # type: ignore[assignment]
            data = [
                [Paragraph(_inline_pdf(cell, force_bold=row_index == 0), body) for cell in row]
                for row_index, row in enumerate(rows)
            ]
            column_count = max(len(row) for row in rows)
            column_width = (A4[0] - 32 * mm) / column_count
            table = Table(data, colWidths=[column_width] * column_count, repeatRows=1, hAlign="LEFT")
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#ECFDF5")),
                        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#C7D9D1")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 6),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                        ("TOPPADDING", (0, 0), (-1, -1), 5),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                    ]
                )
            )
            story.extend([table, Spacer(1, 8)])
    doc.build(story)
    return output.getvalue()
