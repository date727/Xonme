"""Regression tests for native report downloads."""

from io import BytesIO
import unittest
from unittest.mock import patch

from docx import Document
from docx.oxml.ns import qn

from app import report_export


SAMPLE_REPORT = """# 执行摘要 Summary

这是正文 body text，风险等级为 **Critical**。
"""


def _run_fonts(run) -> dict[str, str]:
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    return {
        "ascii": fonts.get(qn("w:ascii")),
        "hAnsi": fonts.get(qn("w:hAnsi")),
        "eastAsia": fonts.get(qn("w:eastAsia")),
    }


class ReportExportTests(unittest.TestCase):
    def test_docx_uses_requested_chinese_and_latin_fonts(self) -> None:
        document = Document(BytesIO(report_export.build_docx(SAMPLE_REPORT)))

        title = document.paragraphs[0]
        heading = document.paragraphs[1]
        body = document.paragraphs[2]

        for paragraph in (title, heading):
            run = paragraph.runs[0]
            self.assertEqual(
                _run_fonts(run),
                {
                    "ascii": "Times New Roman",
                    "hAnsi": "Times New Roman",
                    "eastAsia": "Microsoft YaHei",
                },
            )
            self.assertAlmostEqual(run.font.size.pt, 14)
            self.assertTrue(run.bold)
            self.assertAlmostEqual(paragraph.paragraph_format.line_spacing, 1.5)

        body_run = body.runs[0]
        self.assertEqual(
            _run_fonts(body_run),
            {
                "ascii": "Times New Roman",
                "hAnsi": "Times New Roman",
                "eastAsia": "SimSun",
            },
        )
        self.assertAlmostEqual(body_run.font.size.pt, 12)
        self.assertFalse(body_run.bold)
        self.assertAlmostEqual(body.paragraph_format.line_spacing, 1.5)

    def test_pdf_inline_text_switches_cjk_and_latin_fonts(self) -> None:
        rendered = report_export._inline_pdf("风险 Risk **Critical 严重**")

        self.assertIn('<font name="C2SourceHanSansRegular">风险</font>', rendered)
        self.assertIn('<font name="C2LiberationSerifRegular"> Risk </font>', rendered)
        self.assertIn('<font name="C2LiberationSerifBold">Critical </font>', rendered)
        self.assertIn('<font name="C2SourceHanSansBold">严重</font>', rendered)

    def test_pdf_reports_all_missing_project_fonts(self) -> None:
        missing_directory = report_export.FONT_DIR / "__missing_test_fonts__"
        with patch.object(report_export, "FONT_DIR", missing_directory):
            with self.assertRaises(report_export.ReportFontError) as raised:
                report_export.build_pdf(SAMPLE_REPORT)

        message = str(raised.exception)
        for filename in report_export.PDF_FONT_FILES.values():
            self.assertIn(filename, message)
        self.assertIn(str(missing_directory), message)


if __name__ == "__main__":
    unittest.main()
